"""
obtener-sk-localidades.py — Fuentes oficiales del directorio de localidades de Saskatchewan
=========================================================================================
Solo lectura. Se ejecuta en GitHub Actions (el entorno de desarrollo no
tiene acceso a estos servidores). Guarda los datos casi sin procesar en
scripts/canada-en-datos/sk-fuentes.json (la selección y el cruce se hacen
en scripts/canada-en-datos/compactar-sk.py y en
scripts/canada-en-datos/provincias/saskatchewan.mjs, auditables) y
los imprime comprimidos en el registro entre DATOS-INICIO y DATOS-FIN.

Fuentes:
  - Statistics Canada, límites cartográficos del Censo 2021 (servicio
    ArcGIS de geo.statcan.gc.ca): provincia, divisiones censales,
    regiones económicas, subdivisiones censales (CSD) con su tipo
    oficial, lugares designados (DPL) y aglomeraciones censales (CA).
    Se proyectan a la cónica conforme de Lambert de Statistics Canada
    (EPSG:3347), la misma del mapa de Columbia Británica.
  - Statistics Canada 98-10-0002-01 (Censo 2021: población, viviendas y
    superficie por CSD y CD) y 98-10-0012-01 (lugares designados).
  - Recursos Naturales Canadá, Base de datos de nombres geográficos de
    Canadá (CGNDB), archivo de Saskatchewan: nombres oficiales de lugares poblados.
  - Statistics Canada 14-10-0393-01 y 14-10-0392-01: Encuesta de
    Población Activa por región económica (anual).
  - CMHC, Rental Market Survey 2025 (Saskatchewan, Regina y Saskatoon): tablas 1.1.1 y 1.1.2.
  - Statistics Canada, bases de datos abiertas de centros de salud (ODHF,
    2020) y de establecimientos educativos (ODEF v2.1, 2022): filas de Saskatchewan.

No se rellena ningún dato: lo que falta queda en null.
"""
import base64, csv, datetime, gzip, io, json, re, urllib.parse, urllib.request, zipfile

import openpyxl
from pyproj import Transformer
from shapely.geometry import shape, Point
from shapely.strtree import STRtree

UA = {"User-Agent": "Mozilla/5.0 (verificacion de datos; podcastdelmigrante.com)", "Accept-Language": "en-CA"}
GEO = "https://geo.statcan.gc.ca/geo_wa/rest/services/2021/Cartographic_boundary_files/MapServer"
CGN = "https://ftp.cartes.canada.ca/pub/nrcan_rncan/vector/geobase_cgn_toponyme/prov_csv_eng/cgn_sk_csv_eng.zip"
RMS = "https://assets.cmhc-schl.gc.ca/sites/cmhc/professional/housing-markets-data-and-research/housing-data-tables/rental-market/rental-market-report-data-tables/2025/rmr-{}-2025-en.xlsx"
ODHF = "https://www150.statcan.gc.ca/n1/pub/13-26-0001/2020001/ODHF.zip"
ODEF = "https://www150.statcan.gc.ca/n1/pub/37-26-0001/2022001/ODEF_v2.1.zip"
a_lonlat = Transformer.from_crs("EPSG:3347", "EPSG:4326", always_xy=True)
a_lambert = Transformer.from_crs("EPSG:4326", "EPSG:3347", always_xy=True)
salida = {"consultado": datetime.date.today().isoformat(), "fuentes": {}}


def get(url, t=600):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=t) as r:
        return r.read()


def paso(nombre):
    print(f"\n=== {nombre} ===", flush=True)


def redondea(geo, dec=0):
    def r(c):
        return [r(x) for x in c] if isinstance(c[0], list) else [round(c[0], dec), round(c[1], dec)]
    return {"type": geo["type"], "coordinates": r(geo["coordinates"])}


def capa(n, where="PRUID='47'", offset=None):
    """Capa del servicio de límites del Censo 2021, en EPSG:3347."""
    feats, desde = [], 0
    while True:
        q = {"where": where, "outFields": "*", "returnGeometry": "true", "outSR": "3347", "f": "geojson", "resultOffset": desde, "resultRecordCount": 500}
        if offset:
            q["maxAllowableOffset"] = offset
        d = json.loads(get(f"{GEO}/{n}/query?" + urllib.parse.urlencode(q), 900))
        lote = d.get("features", [])
        feats += lote
        if len(lote) < 500:
            break
        desde += 500
    return feats


def punto(g):
    p = g.representative_point()
    lon, lat = a_lonlat.transform(p.x, p.y)
    return {"x": round(p.x), "y": round(p.y), "lon": round(lon, 5), "lat": round(lat, 5)}


# ------------------------------------------------------------ Límites del Censo 2021
salida["fuentes"]["limites"] = "https://www12.statcan.gc.ca/census-recensement/2021/geo/sip-pis/boundary-limites/index2021-eng.cfm?year=21"

paso("Provincia (contorno del mapa)")
pr = capa(0, offset=150)
salida["provincia_geo"] = [redondea(json.loads(json.dumps(shape(f["geometry"]).simplify(150, preserve_topology=True).__geo_interface__))) for f in pr]
print("polígonos:", len(salida["provincia_geo"]), "propiedades:", pr[0]["properties"])

paso("Divisiones censales")
cd = capa(4, offset=100)
salida["cd"] = []
for f in cd:
    g = shape(f["geometry"]).simplify(150, preserve_topology=True)
    p = f["properties"]
    salida["cd"].append({"CDUID": p["CDUID"], "CDNAME": p["CDNAME"], "CDTYPE": p["CDTYPE"], "LANDAREA": p.get("LANDAREA"), "geo": redondea(json.loads(json.dumps(g.__geo_interface__)))})
print("CD:", [(c["CDUID"], c["CDNAME"], c["CDTYPE"]) for c in salida["cd"]], "bytes:", len(json.dumps(salida["cd"])))

paso("Regiones económicas")
er = capa(2, offset=500)
salida["er"] = {f["properties"]["ERUID"]: f["properties"]["ERNAME"] for f in er}
geom_er = [(f["properties"]["ERUID"], shape(f["geometry"])) for f in er]
def er_de(g):
    p = g.representative_point()
    for k, ge in geom_er:
        if ge.contains(p):
            return k
    return min(geom_er, key=lambda kg: kg[1].distance(p))[0]
for c, f in zip(salida["cd"], cd):
    c["ERUID"] = er_de(shape(f["geometry"]))
print("ER:", salida["er"])
print("CD→ER:", [(c["CDUID"], c["ERUID"]) for c in salida["cd"]])

paso("Subdivisiones censales (CSD)")
csd = capa(9)
geom_csd = [shape(f["geometry"]) for f in csd]
arbol = STRtree(geom_csd)
salida["csd"] = []
for f, g in zip(csd, geom_csd):
    p = f["properties"]
    salida["csd"].append({"CSDUID": p["CSDUID"], "DGUID": p["DGUID"], "CSDNAME": p["CSDNAME"], "CSDTYPE": p["CSDTYPE"], "LANDAREA": p.get("LANDAREA"), "ERUID": er_de(g), **punto(g)})
tipos = {}
for c in salida["csd"]:
    tipos[c["CSDTYPE"]] = tipos.get(c["CSDTYPE"], 0) + 1
print("CSD:", len(salida["csd"]), "tipos:", tipos, "campos:", list(csd[0]["properties"]))


def csd_de(x, y):
    p = Point(x, y)
    for i in arbol.query(p):
        if geom_csd[i].contains(p):
            return csd[i]["properties"]["CSDUID"], 0
    i = arbol.nearest(p)
    return csd[i]["properties"]["CSDUID"], round(geom_csd[i].distance(p))


paso("Aglomeraciones censales (CA)")
cma = capa(6, offset=100)
salida["cma"] = []
for f in cma:
    g = shape(f["geometry"])
    miembros = [c["CSDUID"] for c, gc in zip(salida["csd"], geom_csd) if g.contains(gc.representative_point())]
    salida["cma"].append({"CMAUID": f["properties"]["CMAUID"], "CMANAME": f["properties"]["CMANAME"], "CMATYPE": f["properties"]["CMATYPE"], "csd": miembros})
print("CMA/CA:", [(c["CMAUID"], c["CMANAME"], c["CMATYPE"], len(c["csd"])) for c in salida["cma"]])

paso("Lugares designados (DPL)")
dpl = capa(8)
salida["dpl"] = []
for f in dpl:
    g = shape(f["geometry"])
    pt = punto(g)
    m, dist = csd_de(pt["x"], pt["y"])
    salida["dpl"].append({"DPLUID": f["properties"]["DPLUID"], "DGUID": f["properties"]["DGUID"], "DPLNAME": f["properties"]["DPLNAME"], "DPLTYPE": f["properties"]["DPLTYPE"], "csd": m, "dist": dist, **pt})
print("DPL:", len(salida["dpl"]), "tipos:", sorted({d["DPLTYPE"] for d in salida["dpl"]}))


# ------------------------------------------------------------ Censo 2021 (tablas)
def tabla_statcan(pid):
    zf = zipfile.ZipFile(io.BytesIO(get(f"https://www150.statcan.gc.ca/n1/tbl/csv/{pid}-eng.zip", 900)))
    nombre = [n for n in zf.namelist() if n.endswith(".csv") and "MetaData" not in n][0]
    return zf.open(nombre)


paso("Censo 2021 (98-10-0002-01): provincia, CD y CSD de Saskatchewan")
censo = {}
with tabla_statcan("98100002") as f:
    r = csv.reader(io.TextIOWrapper(f, encoding="utf-8-sig"))
    cab = next(r)
    print("columnas:", cab)
    for fila in r:
        if len(fila) > 2 and (re.match(r"^2021A0005(47\d{5})$", fila[2]) or re.match(r"^2021A0003(47\d{2})$", fila[2]) or fila[2] == "2021A000247"):
            censo[fila[2]] = {cab[i]: v for i, v in enumerate(fila) if v != "" and i >= 1}
print("filas:", len(censo))
salida["censo"] = censo
salida["fuentes"]["censo"] = "https://www150.statcan.gc.ca/t1/tbl1/en/tv.action?pid=9810000201"

paso("Censo 2021 (98-10-0012-01): lugares designados de Saskatchewan")
censo_dpl = {}
with tabla_statcan("98100012") as f:
    r = csv.reader(io.TextIOWrapper(f, encoding="utf-8-sig"))
    cab = next(r)
    for fila in r:
        if len(fila) > 2 and re.match(r"^2021A0006(47\d+)$", fila[2]):
            censo_dpl[fila[2]] = {cab[i]: v for i, v in enumerate(fila) if v != "" and i >= 1}
print("filas:", len(censo_dpl))
salida["censo_dpl"] = censo_dpl
salida["fuentes"]["censo_dpl"] = "https://www150.statcan.gc.ca/t1/tbl1/en/tv.action?pid=9810001201"

# ------------------------------------------------------------ Nombres geográficos oficiales
paso("CGNDB (Recursos Naturales Canadá): lugares poblados de Saskatchewan")
z = zipfile.ZipFile(io.BytesIO(get(CGN, 300)))
filas = list(csv.DictReader(io.TextIOWrapper(z.open(z.namelist()[0]), encoding="utf-8-sig")))
salida["cgn"] = []
for g in filas:
    if g["Generic Category"] != "Populated Place" and g["Generic Term"] not in ("Indian Reserve",):
        continue
    lon, lat = float(g["Longitude"]), float(g["Latitude"])
    x, y = a_lambert.transform(lon, lat)
    m, dist = csd_de(x, y)
    salida["cgn"].append({"id": g["CGNDB ID"], "n": g["Geographical Name"], "t": g["Generic Term"], "cc": g["Concise Code"], "cat": g["Generic Category"], "loc": g["Location"], "fecha": g["Decision Date"], "csd": m, "dist": dist, "x": round(x), "y": round(y), "lat": round(lat, 5), "lon": round(lon, 5)})
cuenta = {}
for g in salida["cgn"]:
    cuenta[(g["t"], g["cat"])] = cuenta.get((g["t"], g["cat"]), 0) + 1
print("nombres:", len(salida["cgn"]), sorted(cuenta.items(), key=lambda x: -x[1]))
print("fuera de una CSD (más de 0 m):", sum(1 for g in salida["cgn"] if g["dist"] > 0))
salida["fuentes"]["cgn"] = "https://natural-resources.canada.ca/earth-sciences/geography/download-geographical-names-data/9245"

# ------------------------------------------------------------ Encuesta de Población Activa
paso("EPA por región económica (Saskatchewan)")
salida["epa"] = {}
for pid, campo in (("14100393", "Labour force characteristics"), ("14100392", "North American Industry Classification System (NAICS)")):
    with tabla_statcan(pid) as f:
        filas = [x for x in csv.DictReader(io.TextIOWrapper(f, encoding="utf-8-sig")) if "Saskatchewan" in x["GEO"]]
    ult = max(x["REF_DATE"] for x in filas)
    print(pid, "geos:", sorted({x["GEO"] for x in filas}), "último año:", ult)
    for x in filas:
        if x["REF_DATE"] != ult:
            continue
        if pid == "14100393" and "Statistics" in x and x["Statistics"] != "Estimate":
            continue
        reg = salida["epa"].setdefault(x["GEO"], {"anio": ult})
        destino = reg.setdefault("lfs" if pid == "14100393" else "industrias", {})
        destino[x[campo]] = {"v": x["VALUE"], "uom": x["UOM"], "status": x["STATUS"], "escala": x["SCALAR_FACTOR"]}
for k, v in salida["epa"].items():
    print(" ", k, v["anio"], (v.get("lfs") or {}).get("Unemployment rate"))

# ------------------------------------------------------------ CMHC 2025
paso("CMHC Rental Market Survey 2025 (Saskatchewan)")
salida["cmhc"] = {}
for c in ["saskatchewan", "regina", "saskatoon"]:
    url = RMS.format(c)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(get(url, 180)), read_only=True, data_only=True)
    except Exception as e:
        print("ERROR", c, e)
        continue
    hojas = {}
    print(c, "hojas:", wb.sheetnames[:20])
    for h in ("Table 1.1.1", "Table 1.1.2", "Table 1.1.3"):
        if h in wb.sheetnames:
            hojas[h] = [[str(x).strip() for x in fila if x not in (None, "")] for fila in wb[h].iter_rows(values_only=True)]
            hojas[h] = [f for f in hojas[h] if f]
    salida["cmhc"][c] = {"url": url, "hojas": hojas}
    for h, filas_h in hojas.items():
        print(" ", h, filas_h[:40])

# ------------------------------------------------------------ Salud y educación (bases abiertas de StatCan)
def csv_zip(url, filtro):
    z = zipfile.ZipFile(io.BytesIO(get(url, 600)))
    n = [x for x in z.namelist() if x.lower().endswith(".csv")][0]
    datos = z.open(n).read()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = datos.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    filas = list(csv.DictReader(io.StringIO(texto)))
    print("archivo:", n, "columnas:", list(filas[0].keys()))
    return [f for f in filas if filtro(f)]


paso("ODHF 2020: centros de salud de Saskatchewan")
try:
    salud = csv_zip(ODHF, lambda f: (f.get("province") or f.get("Province") or "").lower() in ("sk", "saskatchewan"))
    salida["salud"] = []
    for f in salud:
        try:
            lat, lon = float(f.get("latitude") or f.get("Latitude")), float(f.get("longitude") or f.get("Longitude"))
        except (TypeError, ValueError):
            continue
        x, y = a_lambert.transform(lon, lat)
        m, dist = csd_de(x, y)
        salida["salud"].append({**{k: v for k, v in f.items() if v}, "csd": m, "dist": dist, "lat": round(lat, 5), "lon": round(lon, 5)})
    print("centros:", len(salida["salud"]), "tipos:", sorted({(s.get("odhf_facility_type") or s.get("facility_type") or "?") for s in salida["salud"]}))
    salida["fuentes"]["salud"] = "https://www.statcan.gc.ca/en/lode/databases/odhf"
except Exception as e:
    print("ERROR ODHF", type(e).__name__, e)

paso("ODEF v2.1 (2022): establecimientos educativos de Saskatchewan")
try:
    edu = csv_zip(ODEF, lambda f: (f.get("Prov_Terr") or f.get("province") or "").upper() in ("SK", "SASKATCHEWAN"))
    salida["educacion"] = []
    for f in edu:
        try:
            lat, lon = float(f.get("Latitude")), float(f.get("Longitude"))
        except (TypeError, ValueError):
            lat = lon = None
        reg = {k: v for k, v in f.items() if v and v != ".."}
        if lat is not None:
            x, y = a_lambert.transform(lon, lat)
            reg["csd"], reg["dist"] = csd_de(x, y)
        salida["educacion"].append(reg)
    print("establecimientos:", len(salida["educacion"]), "ejemplo:", salida["educacion"][:2])
    salida["fuentes"]["educacion"] = "https://www.statcan.gc.ca/en/lode/databases/odef"
except Exception as e:
    print("ERROR ODEF", type(e).__name__, e)

# ------------------------------------------------------------ Salida
texto = json.dumps(salida, ensure_ascii=False, separators=(",", ":"))
open("scripts/canada-en-datos/sk-fuentes.json", "w", encoding="utf-8").write(texto)
b64 = base64.b64encode(gzip.compress(texto.encode("utf-8"), 9)).decode()
print("\nbytes JSON:", len(texto), "base64:", len(b64))
print("DATOS-INICIO")
for i in range(0, len(b64), 3000):
    print("D|" + b64[i : i + 3000])
print("DATOS-FIN")
