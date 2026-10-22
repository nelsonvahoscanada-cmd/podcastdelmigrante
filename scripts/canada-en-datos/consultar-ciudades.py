"""
consultar-ciudades.py — Datos oficiales para las fichas de las seis
primeras ciudades (Calgary, Edmonton, Brooks, Red Deer, Lethbridge y
Medicine Hat). Solo lectura; imprime lo encontrado para verificarlo.
"""
import csv, io, json, re, urllib.parse, urllib.request, zipfile

import openpyxl

UA = {"User-Agent": "Mozilla/5.0 (verificacion de datos; podcastdelmigrante.com)"}
CIUDADES = ["Calgary", "Edmonton", "Brooks", "Red Deer", "Lethbridge", "Medicine Hat"]
ER = {"4810", "4830", "4850", "4860"}


def get(url, t=300):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=t) as r:
        return r.read(), r.geturl(), r.status


def seccion(t):
    print("\n" + "=" * 8, t, "=" * 8)


# 2. Censo 2021 — tabla 98-10-0002-01 (población por subdivisión de recensement)
seccion("Statistics Canada 98-10-0002-01 (Censo 2021)")
try:
    z, _, _ = get("https://www150.statcan.gc.ca/n1/tbl/csv/98100002-eng.zip", 600)
    zf = zipfile.ZipFile(io.BytesIO(z))
    nombre = [n for n in zf.namelist() if not n.endswith("_MetaData.csv")][0]
    with zf.open(nombre) as f:
        r = csv.reader(io.TextIOWrapper(f, encoding="utf-8-sig"))
        cab = next(r)
        print("  columnas:", cab)
        for fila in r:
            geo = fila[1] if len(fila) > 1 else ""
            if re.match(r"^(Calgary|Edmonton|Brooks|Red Deer|Lethbridge|Medicine Hat)\b", geo) and ("Alberta" in geo or "Alta" in geo or fila[2].startswith("2021A000548")):
                print("  ", fila[:9])
except Exception as e:
    print("  ERROR", e)

# 3. Encuesta de Población Activa por región económica (anual)
for tabla, filtro in (("14100393", r"Unemployment rate|Population|Employment rate|Labour force"), ("14100392", r"")):
    seccion(f"Statistics Canada {tabla} (regiones económicas de Alberta)")
    try:
        z, _, _ = get(f"https://www150.statcan.gc.ca/n1/tbl/csv/{tabla}-eng.zip", 600)
        zf = zipfile.ZipFile(io.BytesIO(z))
        nombre = [n for n in zf.namelist() if not n.endswith("_MetaData.csv")][0]
        with zf.open(nombre) as f:
            r = csv.DictReader(io.TextIOWrapper(f, encoding="utf-8-sig"))
            print("  columnas:", r.fieldnames)
            todas = [fila for fila in r if "Alberta" in fila["GEO"]]
            print("  años:", sorted({f["REF_DATE"] for f in todas})[-4:], "geo:", sorted({f["GEO"] for f in todas})[:12])
            ult = max(f["REF_DATE"] for f in todas)
            filas = [f for f in todas if f["REF_DATE"] == ult]
        if tabla == "14100393":
            for fila in filas:
                if re.search(filtro, fila.get("Labour force characteristics", "")) and fila.get("Statistics", "Estimate") in ("Estimate", ""):
                    print("  ", fila["REF_DATE"], fila["GEO"], fila.get("Labour force characteristics"), fila.get("UOM"), fila["VALUE"], fila.get("STATUS"))
        else:
            for fila in filas:
                if True:
                    print("  ", fila["REF_DATE"], fila["GEO"], "|", fila.get("North American Industry Classification System (NAICS)"), "|", fila.get("UOM"), fila["VALUE"], fila.get("STATUS"))
    except Exception as e:
        print("  ERROR", e)

# 4. Escuelas: Alberta Education, «Alberta schools and authorities»
seccion("Alberta Education: escuelas y autoridades")
try:
    x, _, _ = get("https://education.alberta.ca/media/1626669/authority_and_school.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(x), read_only=True, data_only=True)
    print("  hojas:", wb.sheetnames)
    ws = wb[wb.sheetnames[0]]
    filas = list(ws.iter_rows(values_only=True))
    for k, f in enumerate(filas[:6]):
        print("  cab", k, [c for c in f if c not in (None, "")][:30])
except Exception as e:
    print("  ERROR", e)

# 5. Hospitales: listado oficial de instalaciones (AHCIP facility listing)
seccion("Facility listing")
try:
    import subprocess
    d = json.loads(get("https://open.alberta.ca/api/3/action/package_show?id=alberta-health-facility-and-functional-centre-definitions-and-facility-listing")[0])["result"]
    for r in d["resources"]:
        print("  recurso:", r.get("name"), r.get("format"), r.get("url"))
    url = [r["url"] for r in d["resources"] if (r.get("format") or "").upper() == "PDF" and "listing" in r["url"]][0]
    open("/tmp/f.pdf", "wb").write(get(url)[0])
    subprocess.run(["pdftotext", "-layout", "/tmp/f.pdf", "/tmp/f.txt"], check=True)
    t = open("/tmp/f.txt", encoding="utf-8", errors="ignore").read().split("\n")
    print("  líneas:", len(t))
    for l in t[:40]:
        if l.strip(): print("   |", l[:200])
    for l in t:
        if re.search(r"(Calgary|Edmonton|Brooks|Red Deer|Lethbridge|Medicine Hat)", l) and re.search(r"Hospital|Health Centre|Medical Centre", l, re.I):
            print("   H", re.sub(r"\s{2,}", " | ", l.strip())[:220])
except Exception as e:
    print("  ERROR", e)

# 6. Sitios web municipales y recursos oficiales
seccion("Enlaces oficiales")
for u in ["https://www.edmonton.ca/", "https://www.edmonton.ca/ets", "https://www.calgarytransit.com/", "https://www.calgary.ca/transit", "https://www.brooks.ca/p/transit", "https://www.brooks.ca/p/brooks-transit",
          "https://www.jobbank.gc.ca/jobsearch/jobsearch?locationstring=Brooks%2C+AB", "https://alis.alberta.ca/look-for-work/", "https://www.alberta.ca/find-a-school-board", "https://www.albertahealthservices.ca/findhealth/Search.aspx?q=hospital&loc=Brooks",
          "https://www.alberta.ca/settlement-integration-language-programs", "https://ab.211.ca/", "https://www.alberta.ca/municipal-census", "https://www.alberta.ca/municipal-affairs-population-list"]:
    try:
        b, final, st = get(u, 60)
        t = re.search(rb"<title[^>]*>(.*?)</title>", b, re.S)
        print("  ", st, u, "->", final, "|", (t.group(1).decode("utf-8", "ignore").strip() if t else "")[:90])
    except Exception as e:
        print("  ERR", u, e)
