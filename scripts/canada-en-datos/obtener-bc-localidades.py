"""
obtener-bc-localidades.py — Fuentes oficiales del directorio de localidades de Columbia Británica
================================================================================================
Solo lectura. Se ejecuta en GitHub Actions (el entorno de desarrollo no
tiene acceso a estos servidores). Guarda los datos casi sin procesar en
scripts/canada-en-datos/bc-fuentes.json (la selección y el cruce se hacen
en scripts/canada-en-datos/provincias/british-columbia.mjs, auditables) y
los imprime comprimidos en el registro entre DATOS-INICIO y DATOS-FIN.

Fuentes:
  - Límites y perfil del Censo 2021 de Statistics Canada, en la copia
    oficial del Gobierno de Columbia Británica (BC Geographic Warehouse):
    subdivisiones censales (CSD) con su tipo oficial, distrito regional,
    región económica y área metropolitana; perfil general de cada CSD;
    divisiones censales y regiones económicas. Se proyectan a la cónica
    conforme de Lambert de Statistics Canada (EPSG:3347).
    (El servicio ArcGIS de geo.statcan.gc.ca devolvía error 500.)
  - Statistics Canada, Archivo de atributos geográficos 2021 (92-151-X),
    copia de BC: une cada lugar designado (DPL) con su subdivisión censal.
  - Statistics Canada 98-10-0002-01 (Censo 2021: población por CSD y CD)
    y 98-10-0012-01 (Censo 2021: población de los lugares designados).
  - Gobierno de Columbia Británica, BC Geographical Names (BCGNIS):
    comunidades, localidades y otros lugares poblados con nombre oficial.
  - Gobierno de Columbia Británica, límites legales (ABMS): municipios,
    distritos regionales y áreas electorales vigentes (para contrastar
    el Censo 2021 con la situación actual).
  - Statistics Canada 14-10-0393-01 y 14-10-0392-01: Encuesta de
    Población Activa por región económica (anual).
  - CMHC, Rental Market Survey 2025: tablas 1.1.1 y 1.1.2.
  - Gobierno de Columbia Británica: hospitales (Hospitals in BC) y
    escuelas K-12 (BC Schools K-12).

No se rellena ningún dato: lo que falta queda en null.
"""
import base64, csv, datetime, gzip, io, json, re, urllib.parse, urllib.request, zipfile

import openpyxl
from pyproj import Transformer
from shapely.geometry import shape, Point
from shapely.strtree import STRtree

UA = {"User-Agent": "Mozilla/5.0 (verificacion de datos; podcastdelmigrante.com)", "Accept-Language": "en-CA"}
GEO = "https://geo.statcan.gc.ca/geo_wa/rest/services/2021/Cartographic_boundary_files/MapServer"
WFS = "https://openmaps.gov.bc.ca/geo/pub/wfs"
RMS = "https://assets.cmhc-schl.gc.ca/sites/cmhc/professional/housing-markets-data-and-research/housing-data-tables/rental-market/rental-market-report-data-tables/2025/rmr-{}-2025-en.xlsx"
CMHC_CENTROS = ["british-columbia", "vancouver", "victoria", "kelowna", "abbotsford-mission", "nanaimo", "kamloops", "chilliwack"]
a_lonlat = Transformer.from_crs("EPSG:3347", "EPSG:4326", always_xy=True)
a_lambert = Transformer.from_crs("EPSG:4326", "EPSG:3347", always_xy=True)
salida = {"consultado": datetime.date.today().isoformat(), "fuentes": {}}


def get(url, t=600):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=t) as r:
        return r.read()


def paso(nombre):
    print(f"\n=== {nombre} ===", flush=True)


def redondea(geo, dec=0):
    def r(c):
        return [r(x) for x in c] if isinstance(c[0], list) else [round(c[0], dec), round(c[1], dec)]
    return {"type": geo["type"], "coordinates": r(geo["coordinates"])}


def capa(n, where="PRUID='59'", offset=None):
    """Capa del servicio de límites del Censo 2021, en EPSG:3347."""
    feats, desde = [], 0
    while True:
        q = {"where": where, "outFields": "*", "returnGeometry": "true", "outSR": "3347", "f": "geojson", "resultOffset": desde, "resultRecordCount": 500}
        if offset:
            q["maxAllowableOffset"] = offset
        d = json.loads(get(f"{GEO}/{n}/query?" + urllib.parse.urlencode(q), 900))
        lote = d.get("features", [])
        feats += lote
        if len(lote) < 500:
            break
        desde += 500
    return feats


def punto(geo):
    """Punto representativo (dentro del polígono) en metros y en grados."""
    p = shape(geo).representative_point()
    lon, lat = a_lonlat.transform(p.x, p.y)
    return {"x": round(p.x), "y": round(p.y), "lon": round(lon, 5), "lat": round(lat, 5)}


# ------------------------------------------------------------ Gobierno de BC (WFS)
def wfs(capa_, filtro=None, props=None, srs="EPSG:4326", orden="OBJECTID"):
    feats, desde = [], 0
    while True:
        q = {"service": "WFS", "version": "2.0.0", "request": "GetFeature", "typeName": capa_, "outputFormat": "application/json", "srsName": srs, "count": 5000, "startIndex": desde, "sortBy": orden}
        if filtro:
            q["CQL_FILTER"] = filtro
        if props:
            q["propertyName"] = props
        d = json.loads(get(WFS + "?" + urllib.parse.urlencode(q), 900))
        lote = d.get("features", [])
        feats += lote
        if len(lote) < 5000:
            break
        desde += 5000
    return feats


def a_metros(geo):
    """GeoJSON en grados → geometría shapely en EPSG:3347 (Lambert de Statistics Canada)."""
    from shapely.ops import transform
    return transform(lambda x, y, z=None: a_lambert.transform(x, y), shape(geo))


def punto(g):
    p = g.representative_point()
    lon, lat = a_lonlat.transform(p.x, p.y)
    return {"x": round(p.x), "y": round(p.y), "lon": round(lon, 5), "lat": round(lat, 5)}


# ------------------------------------------------------------ Límites del Censo 2021 (copia del Gobierno de BC)
paso("Subdivisiones censales (CSD), Censo 2021 — CEN_CENSUS_SUBDIVSIONS_SVW")
csd = wfs("WHSE_HUMAN_CULTURAL_ECONOMIC.CEN_CENSUS_SUBDIVSIONS_SVW")
print("CSD:", len(csd), "campos:", list(csd[0]["properties"]))
geom_csd = [a_metros(f["geometry"]) for f in csd]
arbol = STRtree(geom_csd)
CAMPOS_CSD = ("CENSUS_YEAR", "CENSUS_SUBDIVISION_ID", "CENSUS_SUBDIVISION_NAME", "CENSUS_SUBDIVISION_TYPE_CODE", "CENSUS_SUBDIVISION_TYPE_DESC", "CENSUS_DIVISION_ID", "CENSUS_DIVISION_NAME", "CENSUS_DIVISION_TYPE_CODE", "CENSUS_METRO_AREA_ID", "CENSUS_METRO_AREA_NAME", "STAT_AREA_CLASS_TYPE_CODE", "ECONOMIC_REGION_ID", "ECONOMIC_REGION_NAME")
salida["csd"] = [{**{k: f["properties"].get(k) for k in CAMPOS_CSD}, **punto(g)} for f, g in zip(csd, geom_csd)]
tipos = {}
for c in salida["csd"]:
    k = (c["CENSUS_SUBDIVISION_TYPE_CODE"], c["CENSUS_SUBDIVISION_TYPE_DESC"])
    tipos[k] = tipos.get(k, 0) + 1
print("tipos:", sorted(tipos.items(), key=lambda x: -x[1]))
salida["fuentes"]["limites"] = "https://catalogue.data.gov.bc.ca/dataset/current-census-subdivision-boundaries"


def csd_de(x, y):
    p = Point(x, y)
    for i in arbol.query(p):
        if geom_csd[i].contains(p):
            return csd[i]["properties"]["CENSUS_SUBDIVISION_ID"]
    i = arbol.nearest(p)
    return csd[i]["properties"]["CENSUS_SUBDIVISION_ID"]


paso("Perfil del Censo 2021 por CSD — CEN_PROF_CSD_GENERAL_SV")
prof = wfs("WHSE_HUMAN_CULTURAL_ECONOMIC.CEN_PROF_CSD_GENERAL_SV")
print("filas:", len(prof), "campos:", list(prof[0]["properties"]))
salida["perfil"] = {str(f["properties"]["CENSUS_SUBDIVISION_ID"]): {k: v for k, v in f["properties"].items() if k not in ("SE_ANNO_CAD_DATA", "FEATURE_AREA_SQM", "FEATURE_LENGTH_M", "OBJECTID", "CEN_CPCPS_SYSID")} for f in prof if f["properties"].get("CENSUS_YEAR") == 2021}
print("perfiles 2021:", len(salida["perfil"]))
salida["fuentes"]["perfil"] = "https://catalogue.data.gov.bc.ca/dataset/census-profiles-for-bc-census-subdivisions-2021-census"

paso("Divisiones censales (distritos regionales) — CEN_CENSUS_DIVISIONS_SVW")
cd = wfs("WHSE_HUMAN_CULTURAL_ECONOMIC.CEN_CENSUS_DIVISIONS_SVW")
print("CD:", [(f["properties"]["CENSUS_DIVISION_ID"], f["properties"]["CENSUS_DIVISION_NAME"], f["properties"]["CENSUS_DIVISION_TYPE_CODE"], f["properties"]["CENSUS_YEAR"]) for f in cd])
cd = [f for f in cd if f["properties"]["CENSUS_YEAR"] == 2021]
salida["cd"] = []
for f in cd:
    g = a_metros(f["geometry"]).simplify(600, preserve_topology=True)
    salida["cd"].append({"CDUID": f["properties"]["CENSUS_DIVISION_ID"], "CDNAME": f["properties"]["CENSUS_DIVISION_NAME"], "CDTYPE": f["properties"]["CENSUS_DIVISION_TYPE_CODE"], "geo": redondea(json.loads(json.dumps(g.__geo_interface__)))})
print("CD 2021:", len(salida["cd"]), "bytes geometría:", len(json.dumps([c["geo"] for c in salida["cd"]])))

paso("Regiones económicas — CEN_ECONOMIC_REGIONS_SVW")
er = wfs("WHSE_HUMAN_CULTURAL_ECONOMIC.CEN_ECONOMIC_REGIONS_SVW", props="CENSUS_YEAR,ECONOMIC_REGION_ID,ECONOMIC_REGION_NAME,OBJECTID")
salida["er"] = {f["properties"]["ECONOMIC_REGION_ID"]: f["properties"]["ECONOMIC_REGION_NAME"] for f in er if f["properties"]["CENSUS_YEAR"] == 2021}
print("ER:", salida["er"])

# ------------------------------------------------------------ Lugares designados (Archivo de atributos geográficos 2021)
paso("Archivo de atributos geográficos 2021 (92-151-X), BC")
GAF = "https://catalogue.data.gov.bc.ca/dataset/4ed0c389-43ed-4b00-b1be-7c259f47b6a0/resource/6c8eeb86-bb56-48e8-b734-0d25d6a51a7b/download/2021_92-151_x_bc.csv"
filas = list(csv.DictReader(io.StringIO(get(GAF, 900).decode("latin-1"))))
print("bloques:", len(filas), "campos:", list(filas[0]))
print("ejemplo:", filas[0])
def campo(f, *nombres):
    for n in nombres:
        for k in f:
            if k.strip().upper() == n:
                return f[k]
    return None
dpls = {}
for f in filas:
    uid = campo(f, "DPLDGUID_LDIDUGD")
    if not uid or not uid.strip():
        continue
    pob = int(campo(f, "DBPOP2021_IDPOP2021") or 0)
    d = dpls.setdefault(uid.strip(), {"n": campo(f, "DPLNAME_LDNOM"), "tipo": campo(f, "DPLTYPE_LDGENRE"), "pob_db": 0, "csd": {}, "sx": 0.0, "sy": 0.0, "k": 0})
    c = campo(f, "CSDUID_SDRIDU")
    d["csd"][c] = d["csd"].get(c, 0) + pob + 0.001
    d["pob_db"] += pob
    # El archivo no trae coordenadas de bloque: punto representativo de su área de difusión (DA)
    lat, lon = float(campo(f, "DARPLAT_ADLAT") or 0), float(campo(f, "DARPLONG_ADLONG") or 0)
    if lat:
        w = pob + 0.01
        d["sx"] += lon * w
        d["sy"] += lat * w
        d["k"] += w
salida["dpl"] = []
for uid, d in dpls.items():
    lon, lat = d["sx"] / d["k"], d["sy"] / d["k"]
    x, y = a_lambert.transform(lon, lat)
    principal = max(d["csd"].items(), key=lambda kv: kv[1])[0]
    salida["dpl"].append({"uid": uid, "n": d["n"], "tipo": d["tipo"], "lat": round(lat, 5), "lon": round(lon, 5), "x": round(x), "y": round(y), "csd": principal, "csds": sorted(d["csd"]), "pob_db": d["pob_db"]})
print("lugares designados:", len(salida["dpl"]), salida["dpl"][:3])
salida["fuentes"]["gaf"] = "https://catalogue.data.gov.bc.ca/dataset/census-geographic-attribute-file"

# ------------------------------------------------------------ Censo 2021 (tablas)
def tabla_statcan(pid):
    zf = zipfile.ZipFile(io.BytesIO(get(f"https://www150.statcan.gc.ca/n1/tbl/csv/{pid}-eng.zip", 900)))
    nombre = [n for n in zf.namelist() if n.endswith(".csv") and "MetaData" not in n][0]
    return zf.open(nombre)


paso("Censo 2021 (98-10-0002-01): CSD y CD de BC")
censo = {}
with tabla_statcan("98100002") as f:
    r = csv.reader(io.TextIOWrapper(f, encoding="utf-8-sig"))
    cab = next(r)
    print("columnas:", cab)
    for fila in r:
        if len(fila) > 2 and (re.match(r"^2021A0005(59\d{5})$", fila[2]) or re.match(r"^2021A0003(59\d{2})$", fila[2]) or fila[2] == "2021A000259"):
            censo[fila[2]] = {cab[i]: v for i, v in enumerate(fila) if v != "" and i >= 1}
print("filas:", len(censo))
salida["censo"] = censo
salida["censo_columnas"] = cab
salida["fuentes"]["censo"] = "https://www150.statcan.gc.ca/t1/tbl1/en/tv.action?pid=9810000201"

paso("Censo 2021 (98-10-0012-01): lugares designados de BC")
censo_dpl = {}
with tabla_statcan("98100012") as f:
    r = csv.reader(io.TextIOWrapper(f, encoding="utf-8-sig"))
    cab = next(r)
    print("columnas:", cab)
    for fila in r:
        if len(fila) > 2 and re.match(r"^2021A0006(59\d+)$", fila[2]):
            censo_dpl[fila[2]] = {cab[i]: v for i, v in enumerate(fila) if v != "" and i >= 1}
print("filas:", len(censo_dpl))
salida["censo_dpl"] = censo_dpl
salida["fuentes"]["censo_dpl"] = "https://www150.statcan.gc.ca/t1/tbl1/en/tv.action?pid=9810001201"

paso("BC Geographical Names (BCGNIS): lugares poblados")
TIPOS_GN = ["City", "Town", "Community", "Locality", "Urban Community", "Recreational Community", "First Nation Village", "Settlement", "Indian Government District", "Resort Municipality", "Mountain Resort Municipality", "Former Locality"]
filtro = "FEATURE_TYPE IN (" + ",".join(f"'{t}'" for t in TIPOS_GN) + ")"
gn = wfs("WHSE_BASEMAPPING.GNS_GEOGRAPHICAL_NAMES_SP", filtro, props="GEOGRAPHICAL_NAMES_ID,FEATURE_ID,GEOGRAPHICAL_NAME,FEATURE_TYPE,LATITUDE,LONGITUDE,RELATIVE_LOCATION,OBJECTID")
print("lugares:", len(gn))
salida["bcgnis"] = []
for f in gn:
    p = f["properties"]
    if p.get("LATITUDE") is None:
        continue
    x, y = a_lambert.transform(p["LONGITUDE"], p["LATITUDE"])
    salida["bcgnis"].append({"id": p["FEATURE_ID"], "n": p["GEOGRAPHICAL_NAME"], "t": p["FEATURE_TYPE"], "lat": round(p["LATITUDE"], 5), "lon": round(p["LONGITUDE"], 5), "x": round(x), "y": round(y), "rel": p.get("RELATIVE_LOCATION"), "csd": csd_de(x, y)})
cuenta = {}
for g in salida["bcgnis"]:
    cuenta[g["t"]] = cuenta.get(g["t"], 0) + 1
print("por tipo:", cuenta)
salida["fuentes"]["bcgnis"] = "https://catalogue.data.gov.bc.ca/dataset/bc-geographical-names"

paso("Límites legales vigentes (ABMS)")
for clave, capa_, props in (
    ("municipios", "WHSE_LEGAL_ADMIN_BOUNDARIES.ABMS_MUNICIPALITIES_SP", "LGL_ADMIN_AREA_ID,ADMIN_AREA_NAME,ADMIN_AREA_ABBREVIATION,OBJECTID"),
    ("areas_electorales", "WHSE_LEGAL_ADMIN_BOUNDARIES.ABMS_ELECTORAL_AREAS_SP", None),
    ("distritos", "WHSE_LEGAL_ADMIN_BOUNDARIES.ABMS_REGIONAL_DISTRICTS_SP", "LGL_ADMIN_AREA_ID,ADMIN_AREA_NAME,ADMIN_AREA_ABBREVIATION,OBJECTID"),
):
    try:
        fs = wfs(capa_, props=props)
    except Exception as e:
        print("ERROR", clave, e)
        try:
            fs = wfs(capa_)
        except Exception as e2:
            print("ERROR (sin propertyName)", clave, e2)
            continue
    filas = []
    for f in fs:
        p = {k: v for k, v in f["properties"].items() if not isinstance(v, (dict, list)) and k not in ("SE_ANNO_CAD_DATA",)}
        if f.get("geometry"):
            c = shape(f["geometry"]).representative_point()
            p["lat"], p["lon"] = round(c.y, 5), round(c.x, 5)
        filas.append(p)
    salida[clave] = filas
    print(clave, len(filas), "campos:", list(filas[0]) if filas else None)
    for x in filas[:3]:
        print("   ", x)

# ------------------------------------------------------------ Encuesta de Población Activa
paso("EPA por región económica (BC)")
salida["epa"] = {}
for pid, campo in (("14100393", "Labour force characteristics"), ("14100392", "North American Industry Classification System (NAICS)")):
    with tabla_statcan(pid) as f:
        filas = [x for x in csv.DictReader(io.TextIOWrapper(f, encoding="utf-8-sig")) if x["GEO"].endswith(", British Columbia") or x["GEO"] == "British Columbia"]
    ult = max(x["REF_DATE"] for x in filas)
    for x in filas:
        if x["REF_DATE"] != ult:
            continue
        if pid == "14100393" and x.get("Statistics", "Estimate") not in ("Estimate", None) and "Statistics" in x:
            continue
        reg = salida["epa"].setdefault(x["GEO"], {"anio": ult})
        destino = reg.setdefault("lfs" if pid == "14100393" else "industrias", {})
        destino[x[campo]] = {"v": x["VALUE"], "uom": x["UOM"], "status": x["STATUS"], "escala": x["SCALAR_FACTOR"]}
for k, v in salida["epa"].items():
    print(" ", k, v["anio"], (v.get("lfs") or {}).get("Unemployment rate"))

# ------------------------------------------------------------ CMHC 2025
paso("CMHC Rental Market Survey 2025")
salida["cmhc"] = {}
for c in CMHC_CENTROS:
    url = RMS.format(c)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(get(url, 180)), read_only=True, data_only=True)
    except Exception as e:
        print("ERROR", c, e)
        continue
    hojas = {}
    for h in ("Table 1.1.1", "Table 1.1.2"):
        if h in wb.sheetnames:
            hojas[h] = [[str(x).strip() for x in fila if x not in (None, "")] for fila in wb[h].iter_rows(values_only=True)]
            hojas[h] = [f for f in hojas[h] if f]
    salida["cmhc"][c] = {"url": url, "hojas": hojas}
    print(c, "primeras columnas:", [f[0] for f in hojas.get("Table 1.1.2", [])][:60])

# ------------------------------------------------------------ Hospitales y escuelas
paso("Hospitales (Hospitals in BC, CSV de HealthLink BC)")
try:
    d = json.loads(get("https://catalogue.data.gov.bc.ca/api/3/action/package_show?id=hospitals-in-bc", 120))["result"]
    print(d["title"], "| modificado:", d.get("metadata_modified"))
    for r in d["resources"]:
        print("   recurso:", r.get("name"), r.get("format"), r.get("url"))
    rec = next(r for r in d["resources"] if (r.get("format") or "").lower() == "csv")
    filas = list(csv.reader(io.StringIO(get(rec["url"], 300).decode("utf-8-sig", "ignore"))))
    print("   columnas:", filas[0], "filas:", len(filas) - 1)
    print("   ejemplo:", filas[1])
    salida["hospitales"] = {"url": rec["url"], "paquete": "hospitals-in-bc", "modificado": d.get("metadata_modified"), "cabecera": filas[0], "filas": filas[1:]}
except Exception as e:
    print("ERROR hospitales", e)

paso("Hospitales — GSR_HOSPITALS_SVW (con coordenadas)")
try:
    fs = wfs("WHSE_IMAGERY_AND_BASE_MAPS.GSR_HOSPITALS_SVW", orden="SEQUENCE_ID")
    salida["hospitales_wfs"] = [{k: f["properties"].get(k) for k in ("OCCUPANT_NAME", "LOCALITY", "PHYSICAL_ADDRESS", "LATITUDE", "LONGITUDE", "HOURS", "WEBSITE_URL", "DESCRIPTION", "DATE_UPDATED")} for f in fs]
    for h in salida["hospitales_wfs"]:
        x, y = a_lambert.transform(h["LONGITUDE"], h["LATITUDE"])
        h["csd"] = csd_de(x, y)
    print("hospitales:", len(fs))
except Exception as e:
    print("ERROR hospitales WFS", e)

paso("Escuelas K-12 (GSR_SCHOOLS_K_TO_12_SVW)")
try:
    fs = wfs("WHSE_IMAGERY_AND_BASE_MAPS.GSR_SCHOOLS_K_TO_12_SVW", orden="SEQUENCE_ID")
    print("escuelas:", len(fs), "campos:", list(fs[0]["properties"]))
    print("   ejemplo:", fs[0]["properties"])
    utiles = ["OCCUPANT_NAME", "LOCALITY", "DISTRICT_NAME", "DISTRICT_NUMBER", "SCHOOL_CATEGORY", "SCHOOL_EDUCATION_LEVEL", "FACILITY_TYPE", "FRANCOPHONE_PROGRAM_OFFERED", "EARLY_FRENCH_IMMERSION_OFFERED", "LATE_FRENCH_IMMERSION_OFFERED", "DATE_UPDATED"]
    filas_e = []
    for f in fs:
        p = f["properties"]
        x, y = a_lambert.transform(p["LONGITUDE"], p["LATITUDE"]) if p.get("LATITUDE") else (None, None)
        filas_e.append([p.get(k) for k in utiles] + [csd_de(x, y) if x else None])
    salida["escuelas"] = {"campos": utiles + ["csd"], "filas": filas_e}
    cat = {}
    for f in filas_e:
        cat[f[4]] = cat.get(f[4], 0) + 1
    print("   por categoría:", cat)
except Exception as e:
    print("ERROR escuelas", e)

# ------------------------------------------------------------ Salida
texto = json.dumps(salida, ensure_ascii=False, separators=(",", ":"))
open("scripts/canada-en-datos/bc-fuentes.json", "w", encoding="utf-8").write(texto)
b64 = base64.b64encode(gzip.compress(texto.encode("utf-8"), 9)).decode()
print("\nbytes JSON:", len(texto), "base64:", len(b64))
print("DATOS-INICIO")
for i in range(0, len(b64), 3000):
    print("D|" + b64[i : i + 3000])
print("DATOS-FIN")
