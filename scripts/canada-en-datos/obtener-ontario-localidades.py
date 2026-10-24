"""
obtener-ontario-localidades.py — Fuentes oficiales del directorio de localidades de Ontario
=========================================================================================
Solo lectura. Se ejecuta en GitHub Actions (el entorno de desarrollo no
tiene acceso a estos servidores). Guarda los datos casi sin procesar en
scripts/canada-en-datos/ontario-fuentes.json (el cruce y la selección se
hacen en scripts/canada-en-datos/ontario-localidades.mjs, auditables) y
los imprime comprimidos en el registro entre DATOS-INICIO y DATOS-FIN.

Fuentes:
  - Ontario Data Catalogue, «Municipalities» (Ministerio de Asuntos
    Municipales y Vivienda): los 444 municipios con su estatus y nivel
    (single-tier, lower-tier, upper-tier) y su área geográfica.
  - Gobierno de Ontario, Land Information Ontario (LIO): límites
    municipales de nivel inferior y único, y de nivel superior y
    distritos. Se guardan en la proyección cónica conforme de Lambert de
    Statistics Canada (EPSG:3347), la misma del mapa de Canadá.
  - Statistics Canada, «Variant of SGC 2021 for Economic Regions»:
    divisiones censales que forman cada región económica (y excepciones).
  - Statistics Canada 98-10-0002-01: población y viviendas, Censo 2021,
    por subdivisión y división censal.
  - Statistics Canada 14-10-0393-01 y 14-10-0392-01: Encuesta de
    Población Activa por región económica (anual).
  - CMHC, Rental Market Survey 2025: tablas 1.1.1 y 1.1.2 de los centros
    de Ontario.
  - Ontario Data Catalogue: información de las escuelas (si está disponible).

No se rellena ningún dato: lo que falta queda en null.
"""
import base64, csv, datetime, gzip, io, json, re, sys, urllib.request, zipfile
from html.parser import HTMLParser

import openpyxl
from pyproj import Transformer
from shapely.geometry import shape
from shapely.ops import unary_union

UA = {"User-Agent": "Mozilla/5.0 (verificacion de datos; podcastdelmigrante.com)", "Accept-Language": "en-CA"}
RMS = "https://assets.cmhc-schl.gc.ca/sites/cmhc/professional/housing-markets-data-and-research/housing-data-tables/rental-market/rental-market-report-data-tables/2025/rmr-{}-2025-en.xlsx"
CMHC_CENTROS = ["toronto", "ottawa", "hamilton", "london", "kitchener-cambridge-waterloo", "windsor", "oshawa", "st-catharines-niagara", "guelph", "barrie", "kingston", "thunder-bay", "greater-sudbury", "sudbury", "ontario"]
a_lonlat = Transformer.from_crs("EPSG:3347", "EPSG:4326", always_xy=True)
salida = {"consultado": datetime.date.today().isoformat(), "fuentes": {}}


def get(url, t=600):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=t) as r:
        return r.read()


def paso(nombre):
    print(f"\n=== {nombre} ===", flush=True)


# ------------------------------------------------ Lista oficial de municipios
paso("Ontario Data Catalogue: Municipalities")
d = json.loads(get("https://data.ontario.ca/api/3/action/package_show?id=municipalities", 120))["result"]
print(d.get("title"), "| modificado:", d.get("metadata_modified"))
rec = [r for r in d["resources"] if (r.get("format") or "").upper() == "CSV" and re.search(r"_-_en_|english|\ben\b", (r.get("url") or "") + " " + (r.get("name") or ""), re.I)] or [r for r in d["resources"] if (r.get("format") or "").upper() == "CSV"]
for r in d["resources"]:
    print("   recurso:", r.get("name"), r.get("format"), r.get("url"), r.get("last_modified"))
rec = rec[0]
filas = list(csv.reader(io.StringIO(get(rec["url"], 120).decode("utf-8-sig", "ignore"))))
print("columnas:", filas[0], "filas:", len(filas) - 1)
for f in filas[1:6]:
    print("  ", f)
salida["municipios"] = {"paquete": "municipalities", "titulo": d.get("title"), "url": rec["url"], "modificado": rec.get("last_modified") or d.get("metadata_modified"), "cabecera": filas[0], "filas": filas[1:]}
salida["fuentes"]["municipios"] = rec["url"]

# ----------------------------------------- Límites municipales (Land Information Ontario)
LIO = "https://ws.lioservices.lrc.gov.on.ca/arcgis2/rest/services/LIO_OPEN_DATA"
a_3347 = Transformer.from_crs("EPSG:4326", "EPSG:3347", always_xy=True)
from shapely.ops import transform as sh_transform


def redondea(geo):
    """GeoJSON con coordenadas enteras (metros): reduce el tamaño sin perder precisión útil para el mapa."""
    def r(c):
        return [r(x) for x in c] if isinstance(c[0], (list, tuple)) else [round(c[0]), round(c[1])]
    return {"type": geo["type"], "coordinates": r(geo["coordinates"])}


def proyecta(g):
    return sh_transform(lambda x, y, z=None: a_3347.transform(x, y), g)


paso("LIO: capas de límites municipales")
capas = {}
for i in range(1, 11):
    srv = f"{LIO}/LIO_Open{i:02d}/MapServer"
    try:
        info = json.loads(get(srv + "?f=json", 120))
    except Exception as e:
        print("ERROR", srv, e)
        continue
    for l in info.get("layers", []):
        if re.search(r"municipal", l["name"], re.I):
            print(f"  LIO_Open{i:02d}/{l['id']}: {l['name']}")
            capas[l["name"]] = f"{srv}/{l['id']}"


def consulta(url, campos="*", offset=0.0005):
    out, desde = [], 0
    while True:
        q = {"where": "1=1", "outFields": campos, "returnGeometry": "true", "outSR": "4326", "f": "geojson", "maxAllowableOffset": str(offset), "resultOffset": str(desde), "resultRecordCount": "500"}
        fs = json.loads(get(url + "/query?" + __import__("urllib.parse").parse.urlencode(q), 600))["features"]
        out += fs
        if len(fs) < 500:
            return out
        desde += 500


inferior = [u for n, u in capas.items() if re.search(r"lower.*single|single.*lower", n, re.I)]
superior = [u for n, u in capas.items() if re.search(r"upper", n, re.I)]
print("capa inferior/única:", inferior, "| capa superior/distritos:", superior)
salida["lio"] = {"capas": capas}
if inferior:
    fs = consulta(inferior[0])
    print("municipios (inferior/único):", len(fs), "campos:", list(fs[0]["properties"]) if fs else None)
    for f in fs[:3]:
        print("  ", f["properties"])
    salida["lio"]["inferior"] = []
    for f in fs:
        g = proyecta(shape(f["geometry"]))
        p = g.representative_point()
        lon, lat = a_lonlat.transform(p.x, p.y)
        salida["lio"]["inferior"].append({"a": f["properties"], "x": round(p.x), "y": round(p.y), "lat": round(lat, 4), "lon": round(lon, 4), "geo": redondea(g.simplify(1200, preserve_topology=True).__geo_interface__)})
if superior:
    fs = consulta(superior[0])
    print("nivel superior y distritos:", len(fs), "campos:", list(fs[0]["properties"]) if fs else None)
    for f in fs[:3]:
        print("  ", f["properties"])
    salida["lio"]["superior"] = [{"a": f["properties"], "geo": redondea(proyecta(shape(f["geometry"])).simplify(1500, preserve_topology=True).__geo_interface__)} for f in fs]
salida["fuentes"]["limites"] = LIO

# ------------------------------------- Regiones económicas: divisiones censales que las forman
paso("Statistics Canada: variante de la SGC 2021 para regiones económicas")
for url in ("https://www.statcan.gc.ca/en/subjects/standard/sgc/2021/er-additionalinfo",):
    try:
        html = get(url, 120).decode("utf-8", "ignore")
    except Exception as e:
        print("ERROR", url, e)
        continue
    p = Tablas()
    p.feed(html)
    tablas = [[[c["t"] for c in f] for f in t] for t in p.tablas]
    print(url, "tablas:", [len(t) for t in tablas])
    for t in tablas:
        for f in t:
            if any(re.match(r"^35\d{2}$", c) or "Ontario" in c or "Burlington" in c for c in f):
                print("  ", f)
    texto = re.sub(r"<[^>]+>", " ", html)
    notas = [re.sub(r"\s+", " ", m).strip() for m in re.findall(r"[^.]{0,300}(?:Burlington|Halton)[^.]{0,300}\.", texto)]
    for n in notas[:6]:
        print("  nota:", n)
    salida["er_sgc"] = {"url": url, "tablas": tablas, "notas": notas}
    salida["fuentes"]["er"] = url

# ------------------------------------------------------------ Censo 2021
def tabla_statcan(pid):
    zf = zipfile.ZipFile(io.BytesIO(get(f"https://www150.statcan.gc.ca/n1/tbl/csv/{pid}-eng.zip", 900)))
    nombre = [n for n in zf.namelist() if n.endswith(".csv") and "MetaData" not in n][0]
    return zf.open(nombre)


paso("Censo 2021 (98-10-0002-01)")
censo = {}
with tabla_statcan("98100002") as f:
    r = csv.reader(io.TextIOWrapper(f, encoding="utf-8-sig"))
    cab = next(r)
    print("columnas:", cab)
    idx = {c: i for i, c in enumerate(cab)}
    for fila in r:
        if len(fila) < 5:
            continue
        dg = fila[idx["DGUID"]] if "DGUID" in idx else fila[2]
        # CSD de Ontario (A0005 35xxxxx) y divisiones censales (A0003 35xx)
        if re.match(r"^2021A0005(35\d{5})$", dg) or re.match(r"^2021A0003(35\d{2})$", dg):
            censo[dg] = {cab[i]: v for i, v in enumerate(fila) if v != "" and i >= 1}
print("filas de Ontario:", len(censo))
for k in list(censo)[:3]:
    print(" ", k, censo[k])
salida["censo"] = censo
salida["fuentes"]["censo"] = "https://www150.statcan.gc.ca/t1/tbl1/en/tv.action?pid=9810000201"

# ------------------------------------------------ Encuesta de Población Activa
paso("EPA por región económica")
salida["epa"] = {}
for pid, campo in (("14100393", "Labour force characteristics"), ("14100392", "North American Industry Classification System (NAICS)")):
    with tabla_statcan(pid) as f:
        filas = [x for x in csv.DictReader(io.TextIOWrapper(f, encoding="utf-8-sig")) if x["GEO"].endswith(", Ontario") or x["GEO"] == "Ontario"]
    ult = max(x["REF_DATE"] for x in filas)
    for x in filas:
        if x["REF_DATE"] != ult:
            continue
        reg = salida["epa"].setdefault(x["GEO"], {"anio": ult})
        destino = reg.setdefault("lfs" if pid == "14100393" else "industrias", {})
        destino[x[campo]] = {"v": x["VALUE"], "uom": x["UOM"], "status": x["STATUS"], "escala": x["SCALAR_FACTOR"]}
print("EPA regiones:", list(salida["epa"]))
for k, v in salida["epa"].items():
    print(" ", k, v["anio"], (v.get("lfs") or {}).get("Unemployment rate"))

# ------------------------------------------------------------ CMHC 2025
paso("CMHC Rental Market Survey 2025")
salida["cmhc"] = {}
for c in CMHC_CENTROS:
    url = RMS.format(c)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(get(url, 180)), read_only=True, data_only=True)
    except Exception as e:
        print("ERROR", c, e)
        continue
    hojas = {}
    for h in ("Table 1.1.1", "Table 1.1.2"):
        if h in wb.sheetnames:
            hojas[h] = [[str(x).strip() for x in fila if x not in (None, "")] for fila in wb[h].iter_rows(values_only=True)]
            hojas[h] = [f for f in hojas[h] if f]
    salida["cmhc"][c] = {"url": url, "hojas": hojas}
    print(c, "filas:", {h: len(v) for h, v in hojas.items()}, "primeras columnas:", [f[0] for f in hojas.get("Table 1.1.2", [])][:40])

# ------------------------------------------------------------ Escuelas
paso("Escuelas (Ontario Data Catalogue)")
for pkg in ("school-information-and-student-demographics", "ontario-public-schools-contact-information"):
    try:
        d = json.loads(get(f"https://data.ontario.ca/api/3/action/package_show?id={pkg}", 120))["result"]
    except Exception as e:
        print("ERROR", pkg, e)
        continue
    print(pkg, "|", d.get("title"), "| modificado:", d.get("metadata_modified"))
    for r in d["resources"]:
        print("   recurso:", r.get("name"), r.get("format"), r.get("url"), r.get("last_modified"))
    rec = [r for r in d["resources"] if (r.get("format") or "").upper() in ("XLSX", "CSV")]
    if not rec:
        continue
    rec = rec[0]
    try:
        datos = get(rec["url"], 300)
        if rec["url"].lower().endswith(".csv") or (rec.get("format") or "").upper() == "CSV":
            filas = list(csv.reader(io.StringIO(datos.decode("utf-8-sig", "ignore"))))
        else:
            ws = openpyxl.load_workbook(io.BytesIO(datos), read_only=True, data_only=True).worksheets[0]
            filas = [[("" if c is None else str(c)) for c in f] for f in ws.iter_rows(values_only=True)]
        cab = filas[0]
        print("   columnas:", cab)
        print("   ejemplo:", filas[1][:20])
        # Solo las columnas necesarias (consejo, tipo, nombre, nivel, idioma y ciudad)
        utiles = [i for i, c in enumerate(cab) if re.search(r"board|school name|school level|school type|language|city|municipality|special condition", c, re.I)]
        print("   columnas guardadas:", [cab[i] for i in utiles])
        salida["escuelas"] = {"paquete": pkg, "titulo": d.get("title"), "recurso": rec.get("name"), "url": rec["url"], "modificado": rec.get("last_modified") or d.get("metadata_modified"), "cabecera": [cab[i] for i in utiles], "filas": [[f[i] if i < len(f) else "" for i in utiles] for f in filas[1:] if any(f)]}
        break
    except Exception as e:
        print("   ERROR leyendo", e)

# ------------------------------------------------------------ Hospitales
paso("Hospitales (Ontario Data Catalogue)")
for pkg in ("ministry-of-health-service-provider-locations", "general-hospital-locations", "hospital-locations"):
    try:
        d = json.loads(get(f"https://data.ontario.ca/api/3/action/package_show?id={pkg}", 120))["result"]
    except Exception as e:
        print("ERROR", pkg, e)
        continue
    print(pkg, "|", d.get("title"), "| modificado:", d.get("metadata_modified"))
    for r in d["resources"]:
        print("   recurso:", r.get("name"), r.get("format"), r.get("url"))
    rec = [r for r in d["resources"] if (r.get("format") or "").upper() in ("XLSX", "CSV")]
    if not rec:
        continue
    try:
        datos = get(rec[0]["url"], 300)
        if (rec[0].get("format") or "").upper() == "CSV":
            filas = list(csv.reader(io.StringIO(datos.decode("utf-8-sig", "ignore"))))
        else:
            ws = openpyxl.load_workbook(io.BytesIO(datos), read_only=True, data_only=True).worksheets[0]
            filas = [[("" if c is None else str(c)) for c in f] for f in ws.iter_rows(values_only=True)]
        print("   columnas:", filas[0])
        tipos = {}
        for f in filas[1:]:
            for c in f:
                if re.search(r"hospital", c, re.I):
                    tipos[c[:60]] = tipos.get(c[:60], 0) + 1
        print("   valores con 'hospital':", sorted(tipos.items(), key=lambda x: -x[1])[:25])
        salida["hospitales"] = {"paquete": pkg, "titulo": d.get("title"), "url": rec[0]["url"], "modificado": d.get("metadata_modified"), "cabecera": filas[0], "filas": [f for f in filas[1:] if any(re.search(r"hospital", c, re.I) for c in f)]}
        break
    except Exception as e:
        print("   ERROR leyendo", e)

# ------------------------------------------------------------ Salida
texto = json.dumps(salida, ensure_ascii=False, separators=(",", ":"))
open("scripts/canada-en-datos/ontario-fuentes.json", "w", encoding="utf-8").write(texto)
b64 = base64.b64encode(gzip.compress(texto.encode("utf-8"), 9)).decode()
print("\nbytes JSON:", len(texto), "base64:", len(b64))
print("DATOS-INICIO")
for i in range(0, len(b64), 3000):
    print("D|" + b64[i : i + 3000])
print("DATOS-FIN")
