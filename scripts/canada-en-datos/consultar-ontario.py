#!/usr/bin/env python3
"""
consultar-ontario.py — Tablas oficiales de Ontario para «Canadá en Datos»
======================================================================
Herramienta de verificación (no forma parte del sitio).
- Statistics Canada: reutiliza consultar() de consultar-statcan.py con
  filtros de Ontario (población, desempleo, salarios, PIB, censo 2021 de
  la provincia y de las ciudades principales).
- CMHC: tablas del Rental Market Survey 2025 de los centros de Ontario
  (desocupación y alquiler promedio de 2 habitaciones, tablas 1.1.1 y 1.1.2).
Imprime las filas tal como están en los archivos, para auditarlas.
Requiere: pip install openpyxl
Uso: python3 scripts/canada-en-datos/consultar-ontario.py [statcan|cmhc]
"""

import importlib.util
import io
import re
import sys
import urllib.request
from pathlib import Path

import openpyxl

AQUI = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location("statcan", AQUI / "consultar-statcan.py")
statcan = importlib.util.module_from_spec(spec)
spec.loader.exec_module(statcan)

UA = statcan.UA

CIUDADES = r"(Toronto|Ottawa|Hamilton|London|Kitchener|Waterloo|Mississauga|Brampton|Windsor)"

CONSULTAS = [
    ("population", "17100009", {"GEO": r"Ontario"}, 4),
    ("unemploymentRate", "14100287", {"GEO": r"Ontario", "Labour force characteristics": r"(Unemployment rate|Employment rate|Participation rate)", "Gender": r"Total - Gender", "Age group": r"15 years and over", "Statistics": r"Estimate", "Data type": r"Seasonally adjusted"}, 3),
    ("hourlyWage", "14100064", {"GEO": r"Ontario", "Wages": r"(Average|Median) hourly wage rate", "Type of work": r"Both full- and part-time employees", "North American Industry Classification System (NAICS)": r"Total employees, all industries", "Gender": r"Total - Gender", "Age group": r"15 years and over"}, 2),
    ("gdp", "36100222", {"GEO": r"Ontario", "Prices": r"Current prices", "Estimates": r"Gross domestic product at market prices"}, 2),
    ("crimeSeverity", "35100026", {"GEO": r"Ontario( \[\d+\])?", "Statistics": r"Crime severity index"}, 2),
]

CMHC_BASE = "https://assets.cmhc-schl.gc.ca/sites/cmhc/professional/housing-markets-data-and-research/housing-data-tables/rental-market/rental-market-report-data-tables/2025/"
CMHC = ["toronto", "ottawa", "hamilton", "london", "kitchener-cambridge-waterloo", "windsor", "oshawa", "st-catharines-niagara", "guelph", "barrie", "kingston", "sudbury", "thunder-bay"]


def get(url):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=180) as r:
        return r.read(), r.geturl()


def filas(ws, maxfilas=60):
    for fila in ws.iter_rows(min_row=1, max_row=min(ws.max_row, maxfilas), values_only=True):
        celdas = [str(c).strip() for c in fila if c not in (None, "")]
        if celdas:
            yield " | ".join(celdas)


def censo():
    """Tablas del censo 2021 en formato ancho: una fila por región con todas sus columnas."""
    for tabla, geo in (("98100001", r"Ontario"), ("98100002", CIUDADES + r" \((C|CY|CV|T|MU)\), Ontario")):
        print(f"\n=== CENSO — tabla {tabla[:2]}-{tabla[2:4]}-{tabla[4:8]}-01 ===")
        try:
            z = statcan.descargar(tabla)
            nombre = [m for m in z.namelist() if m.endswith(".csv") and "MetaData" not in m][0]
            import csv
            with z.open(nombre) as f:
                lector = csv.reader(io.TextIOWrapper(f, encoding="utf-8-sig"))
                cab = next(lector)
                for fila in lector:
                    if len(fila) > 4 and re.fullmatch(geo, fila[1]):
                        print("  " + fila[1] + " || " + " · ".join(f"{c.split(':')[-1].strip()}={v}" for c, v in zip(cab[4:], fila[4:]) if v and c != "Symbols" and "dwellings" not in c))
        except Exception as e:
            print(f"ERROR censo {tabla}: {e}")


def cmhc_toronto():
    url = f"{CMHC_BASE}rmr-toronto-2025-en.xlsx"
    print(f"\n=== CMHC Toronto CMA (totales) — {url} ===")
    try:
        datos, _ = get(url)
        wb = openpyxl.load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
        for nombre in ("Table 1.1.1", "Table 1.1.2"):
            print(f"--- Hoja {nombre} ---")
            for l in filas(wb[nombre], 120):
                if re.search(r"Toronto CMA|Zone|Oct-2", l) and re.search(r"CMA|Durham|Halton|Peel|York|Oct-2", l):
                    print("  " + l[:260])
    except Exception as e:
        print(f"ERROR Toronto: {e}")


def cmhc():
    for ciudad in CMHC:
        url = f"{CMHC_BASE}rmr-{ciudad}-2025-en.xlsx"
        print(f"\n=== CMHC {ciudad} — {url} ===")
        try:
            datos, _ = get(url)
            wb = openpyxl.load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
            print("Hojas: " + " · ".join(wb.sheetnames[:12]))
            for nombre in wb.sheetnames:
                if re.match(r"^Table 1\.1\.[12]\b", nombre.strip()):
                    print(f"\n--- Hoja {nombre} ---")
                    for l in filas(wb[nombre], 45):
                        print("  " + l[:260])
        except Exception as e:
            print(f"ERROR {ciudad}: {e}")


def main(modos):
    if "delito" in modos:
        statcan.consultar(*CONSULTAS[-1])
    if not modos or "statcan" in modos:
        for c in CONSULTAS:
            try:
                statcan.consultar(*c)
            except Exception as e:
                print(f"ERROR {c[0]}: {e}")
            sys.stdout.flush()
    if not modos or "cmhc" in modos:
        cmhc()
    if not modos or "censo" in modos:
        censo()
    if "toronto" in modos or not modos:
        cmhc_toronto()


if __name__ == "__main__":
    main(sys.argv[1:])
