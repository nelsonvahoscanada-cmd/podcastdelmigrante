"""
obtener-ciudades.py — Datos oficiales de las fichas completas
(Calgary, Edmonton, Brooks, Red Deer, Lethbridge y Medicine Hat).
Solo lectura. Se ejecuta en GitHub Actions; imprime lo verificado y el
JSON resultante comprimido entre DATOS-INICIO y DATOS-FIN.

Fuentes:
  - Statistics Canada 98-10-0002-01: población, Censo 2021.
  - Statistics Canada 14-10-0393-01 y 14-10-0392-01: Encuesta de
    Población Activa por región económica (anual).
  - CMHC, Rental Market Survey 2025 (tablas 1.1.1 y 1.1.2).
  - Alberta Education, «Alberta schools and authorities».
  - Gobierno de Alberta, «Facility and functional centre definitions
    and facility listing» (listado oficial de instalaciones).
"""
import base64, csv, datetime, gzip, io, json, re, urllib.request, zipfile

import openpyxl
from pypdf import PdfReader

UA = {"User-Agent": "Mozilla/5.0 (verificacion de datos; podcastdelmigrante.com)"}
CIUDADES = {"calgary": "Calgary", "edmonton": "Edmonton", "brooks": "Brooks", "red-deer": "Red Deer", "lethbridge": "Lethbridge", "medicine-hat": "Medicine Hat"}
CSD = {"calgary": "2021A00054806016", "edmonton": "2021A00054811061", "brooks": "2021A00054802034", "red-deer": "2021A00054808011", "lethbridge": "2021A00054802012", "medicine-hat": "2021A00054801006"}
RMS = "https://assets.cmhc-schl.gc.ca/sites/cmhc/professional/housing-markets-data-and-research/housing-data-tables/rental-market/rental-market-report-data-tables/2025/rmr-{}-2025-en.xlsx"
CENTRO = {"calgary": ("calgary", "Calgary CMA"), "edmonton": ("edmonton", "Edmonton CMA"), "red-deer": ("alberta", "Red Deer CMA"), "lethbridge": ("alberta", "Lethbridge CMA"), "medicine-hat": ("alberta", "Medicine Hat CA"), "brooks": ("alberta", "Brooks CA")}


def get(url, t=600):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=t) as r:
        return r.read()


def tabla_statcan(pid):
    zf = zipfile.ZipFile(io.BytesIO(get(f"https://www150.statcan.gc.ca/n1/tbl/csv/{pid}-eng.zip")))
    nombre = [n for n in zf.namelist() if not n.endswith("_MetaData.csv")][0]
    return zf.open(nombre)


salida = {"consultado": datetime.date.today().isoformat(), "ciudades": {k: {} for k in CIUDADES}, "regiones": {}}

# ------------------------------------------------------------ Censo 2021
with tabla_statcan("98100002") as f:
    r = csv.reader(io.TextIOWrapper(f, encoding="utf-8-sig"))
    cab = next(r)
    for fila in r:
        for k, dguid in CSD.items():
            if len(fila) > 8 and fila[2] == dguid:
                salida["ciudades"][k]["censo2021"] = {"pob2021": int(fila[4]), "pob2016": int(fila[6]), "cambio": float(fila[8]), "dguid": dguid, "geo": fila[1]}
print("Censo 2021:", {k: v.get("censo2021") for k, v in salida["ciudades"].items()})

# ------------------------------------------------ Encuesta de Población Activa
for pid, campo in (("14100393", "Labour force characteristics"), ("14100392", "North American Industry Classification System (NAICS)")):
    with tabla_statcan(pid) as f:
        filas = [x for x in csv.DictReader(io.TextIOWrapper(f, encoding="utf-8-sig")) if x["GEO"].endswith(", Alberta") or x["GEO"] == "Alberta"]
    ult = max(x["REF_DATE"] for x in filas)
    for x in filas:
        if x["REF_DATE"] != ult:
            continue
        reg = salida["regiones"].setdefault(x["GEO"], {"anio": ult})
        destino = reg.setdefault("lfs" if pid == "14100393" else "industrias", {})
        destino[x[campo]] = {"v": x["VALUE"], "uom": x["UOM"], "status": x["STATUS"], "escala": x["SCALAR_FACTOR"]}
print("EPA regiones:", list(salida["regiones"]))

# ------------------------------------------------------------ CMHC 2025
libros = {}
for k, (archivo, centro) in CENTRO.items():
    if archivo not in libros:
        libros[archivo] = openpyxl.load_workbook(io.BytesIO(get(RMS.format(archivo))), read_only=True, data_only=True)
    wb = libros[archivo]
    res = {"centro": centro, "url": RMS.format(archivo)}
    for hoja, clave in (("Table 1.1.1", "vacancia"), ("Table 1.1.2", "alquiler")):
        for fila in wb[hoja].iter_rows(values_only=True):
            celdas = [str(c).strip() for c in fila if c not in (None, "")]
            if celdas and celdas[0] == centro:
                res[clave] = celdas
    salida["ciudades"][k]["cmhc"] = res
    print("CMHC", k, res.get("alquiler"), res.get("vacancia"))

# ------------------------------------------------------------ Escuelas
x = get("https://education.alberta.ca/media/1626669/authority_and_school.xlsx")
ws = openpyxl.load_workbook(io.BytesIO(x), read_only=True, data_only=True).worksheets[0]
filas = list(ws.iter_rows(values_only=True))
cab = [str(c).strip() if c else "" for c in filas[2]]
print("Escuelas, columnas:", cab)
fecha_informe = None
for fila in filas[3:]:
    d = dict(zip(cab, fila))
    for c in fila:
        if isinstance(c, datetime.datetime):
            fecha_informe = c.date().isoformat()
    ciudad = (d.get("School City") or "").strip()
    for k, nombre in CIUDADES.items():
        if ciudad.lower() == nombre.lower():
            esc = salida["ciudades"][k].setdefault("escuelas", {"total": 0, "por_tipo": {}, "autoridades": {}})
            esc["total"] += 1
            tipo = (d.get("Authority Type") or "").strip()
            esc["por_tipo"][tipo] = esc["por_tipo"].get(tipo, 0) + 1
            if tipo in ("Public", "Separate", "Francophone"):
                aut = (d.get("Authority Name") or "").strip()
                esc["autoridades"][aut] = {"tipo": tipo, "web": (d.get("Authority Website") or "").strip() or None, "n": esc["autoridades"].get(aut, {}).get("n", 0) + 1}
salida["escuelas_fecha"] = fecha_informe
for k in CIUDADES:
    print("Escuelas", k, json.dumps(salida["ciudades"][k].get("escuelas"), ensure_ascii=False)[:600])

# ------------------------------------------------------------ Hospitales
d = json.loads(get("https://open.alberta.ca/api/3/action/package_show?id=alberta-health-facility-and-functional-centre-definitions-and-facility-listing"))["result"]
recurso = [r for r in d["resources"] if not r["name"].startswith("[SUPERSEDED]")][0]
salida["instalaciones"] = {"titulo": recurso["name"], "url": recurso["url"]}
texto = "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(get(recurso["url"]))).pages)
lineas = [l.strip() for l in texto.split("\n") if l.strip()]
print("PDF líneas:", len(lineas))
for l in lineas[:60]:
    print("   |", l[:200])
for l in lineas:
    if re.search(r"\b(Calgary|Edmonton|Brooks|Red Deer|Lethbridge|Medicine Hat)\b", l) and re.search(r"Hospital|Health Centre|Medical Centre|Cancer|Children", l, re.I):
        print("   H", l[:220])
salida["instalaciones"]["lineas"] = [l for l in lineas if re.search(r"\b(Calgary|Edmonton|Brooks|Red Deer|Lethbridge|Medicine Hat)\b", l)]

# ------------------------------------------------- Transporte de Brooks: enlaces del sitio oficial
html = get("https://www.brooks.ca/", 60).decode("utf-8", "ignore")
for m in sorted(set(re.findall(r'href="([^"]*(?:transit|bus|transport)[^"]*)"', html, re.I))):
    print("  brooks enlace:", m)

texto = json.dumps(salida, ensure_ascii=False, separators=(",", ":"))
b64 = base64.b64encode(gzip.compress(texto.encode("utf-8"), 9)).decode()
print("bytes:", len(texto), "base64:", len(b64))
print("DATOS-INICIO")
for i in range(0, len(b64), 3000):
    print("D|" + b64[i : i + 3000])
print("DATOS-FIN")
