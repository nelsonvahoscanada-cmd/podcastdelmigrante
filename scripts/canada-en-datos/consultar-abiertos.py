#!/usr/bin/env python3
"""consultar-abiertos.py — Herramienta temporal de verificación (no forma parte del sitio).
Lee archivos del catálogo oficial de datos abiertos del Gobierno de Alberta
(open.alberta.ca) e imprime las filas y el texto tal como están publicados.
Requiere: pip install openpyxl pypdf"""
import io, json, re, urllib.parse, urllib.request, zipfile
import openpyxl
from pypdf import PdfReader

UA = {"User-Agent": "Mozilla/5.0 (compatible; PodcastDelMigrante-verificacion/1.0)"}
def get(url):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=180) as r:
        return r.read()

def excel(url, ultimas=10, primeras=8):
    print(f"\n=== EXCEL {url}")
    wb = openpyxl.load_workbook(io.BytesIO(get(url)), read_only=True, data_only=True)
    for nombre in wb.sheetnames[:2]:
        ws = wb[nombre]
        filas = [" | ".join(str(c) for c in f if c not in (None, "")) for f in ws.iter_rows(values_only=True)]
        filas = [f for f in filas if f]
        print(f"-- Hoja {nombre} ({len(filas)} filas)")
        for f in filas[:primeras] + ["…"] + filas[-ultimas:]:
            print("  " + f[:240])

def pdf(url, paginas=2, max_chars=5000):
    print(f"\n=== PDF {url}")
    lector = PdfReader(io.BytesIO(get(url)))
    texto = "\n".join((lector.pages[i].extract_text() or "") for i in range(min(paginas, len(lector.pages))))
    print(texto[:max_chars])

def buscar(q, n=5):
    url = "https://open.alberta.ca/api/3/action/package_search?" + urllib.parse.urlencode({"q": q, "rows": n, "sort": "metadata_modified desc"})
    print(f"\n=== BÚSQUEDA open.alberta.ca: {q}")
    j = json.loads(get(url))
    for p in j["result"]["results"]:
        print(f"* {p['title']} · {p['name']} · modificado {p['metadata_modified'][:10]}")
        for r in p.get("resources", [])[:4]:
            print(f"    - {r.get('name')} · {r.get('format')} · {r.get('url')}")

for tarea in [
    lambda: excel("https://open.alberta.ca/dataset/c45f2f78-0d6c-4a7e-98bc-313fbb232040/resource/1b346528-5188-47ee-8fef-b3d41e3db7c0/download/population-estimates-ab-quarterly-1951-to-current.xlsx"),
    lambda: pdf("https://open.alberta.ca/dataset/6b5d7997-6062-474f-a955-837fbab70a89/resource/a4cf065e-5743-400d-bda2-68ec4ba0a808/download/tbf-quarterly-population-report-q2-2026.pdf"),
    lambda: buscar("Alberta economic accounts gross domestic product"),
    lambda: buscar("2021 census population and dwelling counts Alberta"),
    lambda: buscar("crime severity index Alberta"),
]:
    try:
        tarea()
    except Exception as e:
        print(f"ERROR: {e}")

print("\n=== STATISTICS CANADA (nuevo intento)")
for tabla in ("17100009", "35100026", "36100222"):
    try:
        d = get(f"https://www150.statcan.gc.ca/n1/tbl/csv/{tabla}-eng.zip")
        print(f"  {tabla}: {'disponible' if d.startswith(b'PK') else 'no disponible'} ({len(d)//1024} KB)")
    except Exception as e:
        print(f"  {tabla}: ERROR {e}")
