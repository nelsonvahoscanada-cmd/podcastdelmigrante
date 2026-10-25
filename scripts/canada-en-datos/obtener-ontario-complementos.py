"""
obtener-ontario-complementos.py — Fuentes complementarias del directorio de Ontario
===================================================================================
Solo lectura; se ejecuta en GitHub Actions. Completa ontario-fuentes.json
(obtener-ontario-localidades.py) sin volver a descargar los límites:

  - Statistics Canada, Standard Geographical Classification (SGC) 2021:
    divisiones censales que forman cada región económica de Ontario.
  - Ontario Data Catalogue, «School information and student
    demographics»: la edición más reciente en inglés.
  - Hospitales: capas de Land Information Ontario y conjuntos del
    catálogo de Ontario que los listen (si existen).

Escribe scripts/canada-en-datos/ontario-complementos.json y lo imprime
comprimido entre DATOS-INICIO y DATOS-FIN.
"""
import base64, csv, datetime, gzip, io, json, re, urllib.parse, urllib.request, zipfile

import openpyxl

UA = {"User-Agent": "Mozilla/5.0 (verificacion de datos; podcastdelmigrante.com)", "Accept-Language": "en-CA"}
salida = {"consultado": datetime.date.today().isoformat()}


def get(url, t=300):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=t) as r:
        return r.read(), r.geturl()


def texto(html):
    h = re.sub(r"<script[\s\S]*?</script>|<style[\s\S]*?</style>", " ", html)
    h = re.sub(r"<(br|/p|/li|/h\d|/tr|/div|/dt|/dd|/td|/th)[^>]*>", "\n", h, flags=re.I)
    h = re.sub(r"<[^>]+>", " ", h).replace("&nbsp;", " ").replace("&amp;", "&").replace("&ndash;", "–").replace("&#8211;", "–")
    return [re.sub(r"\s+", " ", l).strip() for l in h.split("\n") if l.strip()]


def paso(n):
    print(f"\n=== {n} ===", flush=True)


# ------------------------------------------------ SGC 2021: regiones económicas
paso("SGC 2021: páginas y archivos de la clasificación")
salida["sgc"] = {"paginas": {}, "archivos": {}}
for url in (
    "https://www.statcan.gc.ca/en/subjects/standard/sgc/2021/er-additionalinfo",
    "https://www.statcan.gc.ca/en/subjects/standard/sgc/2021/index",
    "https://www.statcan.gc.ca/en/subjects/standard/sgc/2021/introduction",
):
    try:
        b, final = get(url)
    except Exception as e:
        print("ERROR", url, e)
        continue
    html = b.decode("utf-8", "ignore")
    lineas = texto(html)
    print(url, "→", final, "líneas:", len(lineas))
    for l in lineas:
        if re.search(r"\b35\s?\d{2}\b|Ontario|economic region", l, re.I):
            print("   ·", l[:300])
    enlaces = sorted(set(re.findall(r'href="([^"]+)"', html)))
    utiles = [urllib.parse.urljoin(final, e) for e in enlaces if re.search(r"csv|xlsx|zip|variant|economic|structure|er-|sgc", e, re.I)]
    for e in utiles:
        print("   →", e)
    salida["sgc"]["paginas"][url] = {"lineas": [l for l in lineas if re.search(r"\b35\s?\d{2}\b|Ontario|Halton|Burlington", l)], "enlaces": utiles}
    for e in utiles:
        if re.search(r"\.(csv|xlsx|zip)(\?|$)", e, re.I) and e not in salida["sgc"]["archivos"]:
            try:
                datos, _ = get(e, 300)
            except Exception as ex:
                print("   ERROR", e, ex)
                continue
            print("   archivo", e, len(datos), "bytes")
            filas = []
            if e.lower().endswith(".csv"):
                filas = list(csv.reader(io.StringIO(datos.decode("utf-8-sig", "ignore"))))
            elif e.lower().endswith(".zip"):
                zf = zipfile.ZipFile(io.BytesIO(datos))
                for n in zf.namelist():
                    if n.lower().endswith(".csv"):
                        filas += list(csv.reader(io.StringIO(zf.read(n).decode("utf-8-sig", "ignore"))))
            elif e.lower().endswith(".xlsx"):
                ws = openpyxl.load_workbook(io.BytesIO(datos), read_only=True, data_only=True).worksheets[0]
                filas = [[("" if c is None else str(c)) for c in f] for f in ws.iter_rows(values_only=True)]
            print("   cabecera:", filas[:2])
            on = [f for f in filas if any(re.match(r"^35\d{0,5}$", c.strip()) for c in f) or any("Ontario" in c for c in f)]
            print("   filas de Ontario:", len(on))
            for f in on[:15]:
                print("     ", f)
            salida["sgc"]["archivos"][e] = {"cabecera": filas[0] if filas else [], "filas": on}

# ------------------------------------------------ Escuelas: edición más reciente
paso("Escuelas: edición más reciente")
d = json.loads(get("https://data.ontario.ca/api/3/action/package_show?id=school-information-and-student-demographics")[0])["result"]
en = [r for r in d["resources"] if (r.get("format") or "").upper() == "XLSX" and re.search(r"_en[_.]", r["url"])]
en.sort(key=lambda r: re.findall(r"20\d\d_\d\d", r["url"])[-1] if re.findall(r"20\d\d_\d\d", r["url"]) else "", reverse=True)
for r in en:
    print("  ", r["name"], r["url"], r.get("last_modified"))
rec = en[0]
print("Elegido:", rec["name"], rec["url"])
ws = openpyxl.load_workbook(io.BytesIO(get(rec["url"])[0]), read_only=True, data_only=True).worksheets[0]
filas = [[("" if c is None else str(c)) for c in f] for f in ws.iter_rows(values_only=True)]
cab = filas[0]
utiles = [i for i, c in enumerate(cab) if re.search(r"board name|board type|board website|school name|school type|school level|school language|^municipality$|^city$|special condition", c, re.I)]
print("columnas:", cab)
print("columnas guardadas:", [cab[i] for i in utiles], "filas:", len(filas) - 1)
salida["escuelas"] = {"titulo": d.get("title"), "recurso": rec["name"], "url": rec["url"], "modificado": rec.get("last_modified"), "cabecera": [cab[i] for i in utiles], "filas": [[f[i] if i < len(f) else "" for i in utiles] for f in filas[1:] if any(f)]}

# ------------------------------------------------ Hospitales
paso("Hospitales: capas de Land Information Ontario")
LIO = "https://ws.lioservices.lrc.gov.on.ca/arcgis2/rest/services/LIO_OPEN_DATA"
capas = {}
for i in range(1, 11):
    try:
        info = json.loads(get(f"{LIO}/LIO_Open{i:02d}/MapServer?f=json", 120)[0])
    except Exception as e:
        print("ERROR", i, e)
        continue
    for l in info.get("layers", []):
        if re.search(r"hospital|health|medical", l["name"], re.I):
            print(f"  LIO_Open{i:02d}/{l['id']}: {l['name']}")
            capas[l["name"]] = f"{LIO}/LIO_Open{i:02d}/MapServer/{l['id']}"
salida["hospitales"] = {"capas": capas}
for nombre, url in capas.items():
    if not re.search(r"hospital", nombre, re.I):
        continue
    try:
        q = {"where": "1=1", "outFields": "*", "returnGeometry": "false", "f": "json", "resultRecordCount": "2000"}
        fs = json.loads(get(url + "/query?" + urllib.parse.urlencode(q), 300)[0])["features"]
    except Exception as e:
        print("ERROR", nombre, e)
        continue
    print(nombre, "registros:", len(fs), "campos:", list(fs[0]["attributes"]) if fs else None)
    for f in fs[:5]:
        print("   ", f["attributes"])
    salida["hospitales"][nombre] = {"url": url, "filas": [f["attributes"] for f in fs]}

paso("Hospitales: catálogo de datos de Ontario")
for q in ("hospital locations", "hospitals"):
    r = json.loads(get("https://data.ontario.ca/api/3/action/package_search?rows=8&q=" + urllib.parse.quote(q))[0])["result"]
    for p in r["results"]:
        print("  ", p["name"], "|", p["title"], "|", [(x.get("format"), x.get("url")) for x in p["resources"]][:3])

texto_json = json.dumps(salida, ensure_ascii=False, separators=(",", ":"))
open("scripts/canada-en-datos/ontario-complementos.json", "w", encoding="utf-8").write(texto_json)
b64 = base64.b64encode(gzip.compress(texto_json.encode("utf-8"), 9)).decode()
print("\nbytes JSON:", len(texto_json), "base64:", len(b64))
print("DATOS-INICIO")
for i in range(0, len(b64), 3000):
    print("D|" + b64[i : i + 3000])
print("DATOS-FIN")
