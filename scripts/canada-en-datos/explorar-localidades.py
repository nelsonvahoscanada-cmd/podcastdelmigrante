"""Exploración (solo lectura) de las fuentes oficiales de localidades de Alberta."""
import json, io, re, urllib.request, collections

UA = {"User-Agent": "Mozilla/5.0 (verificacion de datos; podcastdelmigrante.com)"}
def get(url):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=120) as r:
        return r.read()
def js(url): return json.loads(get(url))
G = "https://geospatial.alberta.ca/titan/rest/services"

# 1. Servicios de límites disponibles
for folder in ["boundaries", ""]:
    try:
        d = js(f"{G}/{folder}?f=json")
        print("== carpeta", folder or "(raíz)", "folders:", d.get("folders"), "services:", [s["name"] for s in d.get("services", [])])
    except Exception as e: print("ERR", folder, e)

for folder in ["census", "economics", "society"]:
    try:
        d = js(f"{G}/{folder}?f=json"); print("== carpeta", folder, [s["name"] + ":" + s["type"] for s in d.get("services", [])])
    except Exception as e: print("ERR", folder, e)
for svc in ["boundaries/municipal_metis_settlement_public"]:
    d = js(f"{G}/{svc}/FeatureServer/layers?f=json")
    for l in d.get("layers", []):
        c = js(f"{G}/{svc}/FeatureServer/{l['id']}/query?where=1%3D1&returnCountOnly=true&f=json")
        print(" metis layer", l["id"], l["name"], [f["name"] for f in l["fields"]], c)

# 2. Tipos de la capa de puntos (Hamlet, Locality, Townsite)
d = js(f"{G}/boundaries/municipal_communities_public/FeatureServer/0/query?where=1%3D1&outFields=TYPE&returnGeometry=false&f=json&resultRecordCount=5000")
print("TIPOS punto:", collections.Counter(f["attributes"]["TYPE"] for f in d["features"]), "n=", len(d["features"]), "exceeded", d.get("exceededTransferLimit"))
print("maxRecordCount:", js(f"{G}/boundaries/municipal_communities_public/FeatureServer/0?f=json").get("maxRecordCount"))

# 3. Lista de población 2025 (Municipal Affairs)
import openpyxl
x = get("https://open.alberta.ca/dataset/8dae0ffd-dbb7-46f7-89eb-efa64b2ce385/resource/387c1c81-cfcb-4616-be4e-dc829c7fd6fb/download/2025-mapl-worksheets-for-internet-posting.xlsx")
wb = openpyxl.load_workbook(io.BytesIO(x), data_only=True)
for ws in wb.worksheets:
    print("== hoja", ws.title, ws.max_row, "x", ws.max_column)
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 14 or i in (ws.max_row - 3, ws.max_row - 2, ws.max_row - 1):
            print("  ", i, [c for c in row if c is not None][:14])

# 4. StatCan: clasificación geográfica (división censal → región económica)
for u in ["https://www.statcan.gc.ca/en/subjects/standard/sgc/2021/introduction",
          "https://www150.statcan.gc.ca/n1/en/catalogue/12-571-X",
          "https://www12.statcan.gc.ca/census-recensement/2021/geo/sip-pis/boundary-limites/index2021-eng.cfm",
          "https://www.statcan.gc.ca/en/subjects/standard/sgc/2021/sgc-cgt-2021-structure-eng.csv",
          "https://www150.statcan.gc.ca/t1/tbl1/en/tv.action?pid=9810000201"]:
    try:
        b = get(u)[:300000].decode("utf-8", "ignore")
        t = re.search(r"<title>(.*?)</title>", b, re.S)
        print("OK", u, (t.group(1).strip() if t else b[:200].replace("\n", " "))[:160])
        for m in sorted(set(re.findall(r'href="([^"]+\.(?:csv|zip|xlsx?))"', b)))[:30]: print("    link", m)
    except Exception as e: print("ERR", u, e)
