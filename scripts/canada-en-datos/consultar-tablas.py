#!/usr/bin/env python3
"""
consultar-tablas.py — Lectura de tablas oficiales en Excel para «Canadá en Datos»
======================================================================
Herramienta de verificación (no forma parte del sitio).
- CMHC: tablas del Rental Market Survey 2025 de Calgary y Edmonton
  (tasa de desocupación y alquiler promedio de 2 habitaciones).
- Gobierno de Alberta: tablas de estimaciones de población (Excel enlazado
  desde alberta.ca/demographic-statistics).
- Statistics Canada: nuevo intento con la descarga oficial de tablas.
Imprime las filas tal como están en los archivos, para auditarlas.
Requiere: pip install openpyxl
"""

import io
import re
import urllib.request
import zipfile

import openpyxl

UA = {"User-Agent": "Mozilla/5.0 (compatible; PodcastDelMigrante-verificacion/1.0)"}

CMHC = {
    "Calgary": "https://assets.cmhc-schl.gc.ca/sites/cmhc/professional/housing-markets-data-and-research/housing-data-tables/rental-market/rental-market-report-data-tables/2025/rmr-calgary-2025-en.xlsx",
    "Edmonton": "https://assets.cmhc-schl.gc.ca/sites/cmhc/professional/housing-markets-data-and-research/housing-data-tables/rental-market/rental-market-report-data-tables/2025/rmr-edmonton-2025-en.xlsx",
}


def get(url):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=180) as r:
        return r.read(), r.geturl()


def filas(ws, maxfilas=60):
    for fila in ws.iter_rows(min_row=1, max_row=min(ws.max_row, maxfilas), values_only=True):
        celdas = [str(c).strip() for c in fila if c not in (None, "")]
        if celdas:
            yield " | ".join(celdas)


def cmhc():
    for ciudad, url in CMHC.items():
        print(f"\n=== CMHC {ciudad} — {url} ===")
        try:
            datos, _ = get(url)
            wb = openpyxl.load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
            print("Hojas: " + " · ".join(wb.sheetnames))
            for nombre in wb.sheetnames:
                if re.match(r"^Table 1\.1\.[12]\b", nombre) or nombre.strip() in ("Table 1.1.1", "Table 1.1.2"):
                    print(f"\n--- Hoja {nombre} ---")
                    for l in filas(wb[nombre], 45):
                        print("  " + l[:260])
        except Exception as e:
            print(f"ERROR {ciudad}: {e}")


def poblacion():
    pagina = "https://www.alberta.ca/demographic-statistics"
    print(f"\n=== POBLACIÓN {pagina} ===")
    try:
        html, _ = get(pagina)
        html = html.decode("utf-8", "replace")
        enlaces = sorted(set(re.findall(r'href="([^"]+\.(?:xlsx|xls|pdf)[^"]*)"', html, re.I)))
        for e in enlaces:
            print("  enlace: " + e)
        excel = [e for e in enlaces if re.search(r"\.xlsx?", e, re.I) and re.search(r"estimat|component|quarter", e, re.I)]
        for e in excel[:2]:
            url = e if e.startswith("http") else "https://www.alberta.ca" + e
            print(f"\n--- Excel {url} ---")
            datos, final = get(url)
            print(f"  URL final: {final}")
            wb = openpyxl.load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
            print("  Hojas: " + " · ".join(wb.sheetnames))
            for nombre in wb.sheetnames[:3]:
                ws = wb[nombre]
                print(f"\n  >> Hoja {nombre} (filas {ws.max_row})")
                todas = list(filas(ws, 400))
                for l in todas[:12] + ["…"] + todas[-14:]:
                    print("    " + l[:220])
    except Exception as e:
        print(f"ERROR población: {e}")


def statcan():
    print("\n=== STATISTICS CANADA (nuevo intento) ===")
    for tabla in ("17100009", "35100026", "36100222", "98100001"):
        url = f"https://www150.statcan.gc.ca/n1/tbl/csv/{tabla}-eng.zip"
        try:
            datos, _ = get(url)
            if not datos.startswith(b"PK"):
                t = re.search(rb"<title>([^<]*)</title>", datos)
                print(f"  {tabla}: no disponible — {t.group(1).decode('utf-8', 'replace') if t else datos[:80]}")
                continue
            z = zipfile.ZipFile(io.BytesIO(datos))
            print(f"  {tabla}: disponible ({len(datos) // 1024} KB) {z.namelist()}")
        except Exception as e:
            print(f"  {tabla}: ERROR {e}")


if __name__ == "__main__":
    cmhc()
    poblacion()
    statcan()
